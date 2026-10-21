"""Build the sample CV used for testing the upload flow.

    cd apps/api && uv run python ../../test-data/make_cv.py

Writes a Word file next to this script. The person is invented; the shape is a real student CV,
including the things that trip parsers: a label line, a tool with punctuation, and a CGPA.
"""

import pathlib

from docx import Document
from docx.shared import Pt

OUT = pathlib.Path(__file__).resolve().parent / "ifeoma-chukwu-cv.docx"

LINES: list[tuple[str, bool]] = [
    ("IFEOMA CHUKWU", True),
    ("Lagos, Nigeria | ifeoma.chukwu@gmail.com | 0803 555 0142 | github.com/ifeomac", False),
    ("", False),
    ("EDUCATION", True),
    ("University of Lagos - BSc Computer Science, 2022 - 2026", False),
    ("CGPA: 4.31/5", False),
    ("", False),
    ("EXPERIENCE", True),
    ("Andela - Frontend Engineering Intern, June 2025 - September 2025", False),
    ("- Fixed bugs in the internal staff dashboard built with React and TypeScript", False),
    ("- Added a filter to the reports page so staff could find old reports faster", False),
    ("- Wrote the onboarding notes the next set of interns used", False),
    ("Technologies: React, TypeScript, Next.js, Git", False),
    ("", False),
    ("Department of Computer Science - Teaching Assistant, January 2025 - May 2025", False),
    ("- Ran weekly lab sessions for first-year students learning Python", False),
    ("", False),
    ("PROJECTS", True),
    ("Lecture Reminder Bot - a WhatsApp bot that sends students their timetable each morning", False),
    ("- Built with Python and the Twilio API", False),
    ("- Used by about 60 students in my department last semester", False),
    ("Market Price Tracker - collects prices from local online shops and charts them", False),
    ("- Built with Python and SQLite", False),
    ("", False),
    ("SKILLS", True),
    ("JavaScript, TypeScript, React, Python, SQL, Git", False),
    ("", False),
    ("CERTIFICATIONS", True),
    ("Google IT Support Professional Certificate, 2024", False),
]


def main() -> None:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for text, bold in LINES:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(text)
        run.bold = bold
    document.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()

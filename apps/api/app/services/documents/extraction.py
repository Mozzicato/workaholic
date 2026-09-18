"""Getting text out of an uploaded document (PRD §37, §39).

An upload is untrusted input from outside the system: it may be malformed, enormous, or crafted
to be neither a PDF nor a DOCX whatever its name claims. So the file is identified by its actual
bytes rather than by its name or the browser's Content-Type, parsed inside a guard, and reduced
to plain text. The bytes are never stored and never executed.

What comes out is text, which is still untrusted — it is the author's words, and everything
downstream treats it as data (PRD §38).
"""

import io
import re
from dataclasses import dataclass
from enum import StrEnum

import structlog
from docx import Document as DocxDocument
from pypdf import PdfReader
from pypdf.errors import PyPdfError

from app.core.errors import AppError, ErrorCode

logger = structlog.get_logger(__name__)

__all__ = ["MAX_UPLOAD_BYTES", "ExtractedDocument", "UploadKind", "extract_text", "pdf_text"]

MAX_UPLOAD_BYTES = 8 * 1024 * 1024
"""A resume is a page or two. Anything larger is a mistake or an attack, not a CV."""

MAX_EXTRACTED_CHARACTERS = 60_000
MAX_PDF_PAGES = 30

_PDF_MAGIC = b"%PDF-"
_ZIP_MAGIC = b"PK\x03\x04"  # DOCX is a zip archive
_WHITESPACE = re.compile("[ \t\u00a0]+")  # PDFs are full of no-break spaces
_BLANK_LINES = re.compile(r"\n{3,}")


class UploadKind(StrEnum):
    PDF = "pdf"
    DOCX = "docx"


class UnsupportedDocumentError(AppError):
    def __init__(self, message: str) -> None:
        super().__init__(ErrorCode.VALIDATION_ERROR, message, 422)


class UnreadableDocumentError(AppError):
    def __init__(self) -> None:
        super().__init__(
            ErrorCode.VALIDATION_ERROR,
            "We couldn't read that file. If it is a scan or an image, the text can't be "
            "extracted — type or paste what it says instead.",
            422,
        )


@dataclass(frozen=True, slots=True)
class ExtractedDocument:
    kind: UploadKind
    text: str
    pages: int


def identify(content: bytes) -> UploadKind:
    """What this file actually is, from its bytes. The filename is not evidence."""
    if content.startswith(_PDF_MAGIC):
        return UploadKind.PDF
    if content.startswith(_ZIP_MAGIC):
        return UploadKind.DOCX
    raise UnsupportedDocumentError("That file isn't a PDF or a Word document.")


def extract_text(content: bytes) -> ExtractedDocument:
    """Plain text from an uploaded PDF or DOCX, or a clear refusal."""
    if not content:
        raise UnsupportedDocumentError("That file was empty.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise UnsupportedDocumentError("That file is too large. Resumes are a page or two.")

    kind = identify(content)
    text = _read_pdf(content) if kind is UploadKind.PDF else _read_docx(content)
    cleaned = _tidy(text)
    if len(cleaned) < 40:
        raise UnreadableDocumentError()
    return ExtractedDocument(
        kind=kind, text=cleaned[:MAX_EXTRACTED_CHARACTERS], pages=_page_count(content, kind)
    )


def pdf_text(content: bytes) -> str:
    """Text read back out of a PDF we produced, for the ATS extractability check."""
    try:
        return _read_pdf(content)
    except AppError:
        return ""


def _read_pdf(content: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise UnsupportedDocumentError("That PDF is password-protected.")
        pages = reader.pages[:MAX_PDF_PAGES]
        return "\n".join(page.extract_text() or "" for page in pages)
    except UnsupportedDocumentError:
        raise
    except (PyPdfError, ValueError, OSError) as exc:
        # The parser's message can echo file contents, so only the type is logged.
        logger.warning("document.pdf.unreadable", error=type(exc).__name__)
        raise UnreadableDocumentError() from exc


def _read_docx(content: bytes) -> str:
    try:
        document = DocxDocument(io.BytesIO(content))
    except Exception as exc:  # python-docx raises a wide range on malformed input
        logger.warning("document.docx.unreadable", error=type(exc).__name__)
        raise UnreadableDocumentError() from exc

    parts = [paragraph.text for paragraph in document.paragraphs]
    # Tables are where people hide their whole work history in a Word CV.
    for table in document.tables:
        for row in table.rows:
            parts.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _page_count(content: bytes, kind: UploadKind) -> int:
    if kind is not UploadKind.PDF:
        return 0
    try:
        return len(PdfReader(io.BytesIO(content)).pages)
    except (PyPdfError, ValueError, OSError):  # pragma: no cover - already parsed once
        return 0


def _tidy(text: str) -> str:
    collapsed = _WHITESPACE.sub(" ", text.replace("\r\n", "\n").replace("\r", "\n"))
    lines = [line.strip() for line in collapsed.split("\n")]
    return _BLANK_LINES.sub("\n\n", "\n".join(lines)).strip()
